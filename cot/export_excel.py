"""集計結果をExcel（.xlsx）に出力する。

シート構成：
- 説明                 : 出典・用語・基準日
- サマリー_先物のみ     : 最新週の全銘柄×全区分
- サマリー_先物OP込み   : 同上（先物+オプション）
- 履歴_TFF / 履歴_Disagg / 履歴_Legacy / 履歴_CIT : 全期間の縦持ちデータ
"""

import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from cot.reports import REPORTS, SPECULATOR_CATEGORIES, VARIANT_COMBINED, VARIANT_FUTURES_ONLY, VARIANT_LABELS

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def signal_label(index_value):
    """COTインデックスを表示用の目安ラベルに変換する。"""
    if pd.isna(index_value):
        return ""
    if index_value >= 90:
        return "買い極端"
    if index_value <= 10:
        return "売り極端"
    if index_value >= 80 or index_value <= 20:
        return "要注意"
    return "中立"


def attach_labels(frame, markets):
    """銘柄名・資産区分・レポート名・区分名・並び順を付ける。"""
    labeled = frame.merge(markets[["code", "name_ja", "asset_class", "market_order"]], on="code", how="left")
    labeled["report_label"] = labeled["report"].map(lambda key: REPORTS[key]["label"])
    labeled["report_order"] = labeled["report"].map(lambda key: list(REPORTS).index(key))
    labeled["category_label"] = labeled.apply(
        lambda row: REPORTS[row["report"]]["category_labels"][row["category"]], axis=1
    )
    labeled["category_order"] = labeled.apply(
        lambda row: list(REPORTS[row["report"]]["categories"]).index(row["category"]), axis=1
    )
    labeled["variant_label"] = labeled["variant"].map(VARIANT_LABELS)
    return labeled


def build_summary(latest, variant):
    """サマリーシート用の表を作る。"""
    rows = latest[latest["variant"] == variant].copy()
    rows = rows.sort_values(["report_order", "market_order", "category_order"])
    rows["signal"] = rows["cot_index"].map(signal_label)
    table = pd.DataFrame({
        "基準日": rows["date"],
        "レポート": rows["report_label"],
        "資産区分": rows["asset_class"],
        "銘柄": rows["name_ja"],
        "CFTCコード": rows["code"],
        "投機筋区分": rows["category_label"],
        "ロング": rows["long"],
        "ショート": rows["short"],
        "ネット": rows["net"],
        "前週比": rows["net_change"],
        "総建玉": rows["open_interest"],
        "建玉比(%)": rows["net_oi_pct"],
        "COT指数": rows["cot_index"],
        "シグナル": rows["signal"],
    })
    return table


def build_history(metrics, report_key=None, categories=None):
    """履歴の表を作る。

    report_key と categories を指定すると、そのレポートの指定区分だけに絞る（Excelの履歴シート用）。
    どちらも省略すると全レポート・全区分（CSV用）。
    """
    rows = metrics
    if report_key is not None:
        rows = rows[rows["report"] == report_key]
    if categories is not None:
        rows = rows[rows["category"].isin(categories)]
    rows = rows.sort_values(["report_order", "date", "variant", "market_order", "category_order"])
    table = pd.DataFrame({
        "レポート": rows["report"],
        "基準日": rows["date"],
        "データ種類": rows["variant_label"],
        "資産区分": rows["asset_class"],
        "銘柄": rows["name_ja"],
        "CFTCコード": rows["code"],
        "投機筋区分": rows["category_label"],
        "ロング": rows["long"],
        "ショート": rows["short"],
        "ネット": rows["net"],
        "前週比": rows["net_change"],
        "総建玉": rows["open_interest"],
        "建玉比(%)": rows["net_oi_pct"],
        "COT指数": rows["cot_index"],
    })
    return table


def build_readme(latest_date, generated_at):
    """説明シートの内容を作る。"""
    lines = [
        ["CFTC COT 週次集計"],
        [f"基準日（火曜時点の建玉）: {latest_date}"],
        [f"作成日時: {generated_at}"],
        [""],
        ["出典: U.S. Commodity Futures Trading Commission (CFTC), Commitments of Traders"],
        ["https://publicreporting.cftc.gov/stories/s/r4w3-av2u"],
        [""],
        ["用語"],
        ["ネット = ロング - ショート（スプレッド建玉は含めない）"],
        ["前週比 = 今週ネット - 前週ネット"],
        ["建玉比(%) = ネット ÷ 総建玉 × 100"],
        ["COT指数 = 過去156週のネットの最小～最大レンジ内での位置（0～100）。52週以上のデータがあれば計算"],
        ["シグナル = COT指数 90以上:買い極端 / 10以下:売り極端 / 80以上・20以下:要注意 / それ以外:中立"],
        [""],
        ["履歴シート"],
        ["投機筋区分のみ掲載（TFF:レバレッジドファンド・アセットマネージャー / Disaggregated:マネージドマネー / Legacy:非商業 / CIT:指数トレーダー）"],
        ["全区分の履歴は同じフォルダの「YYYYMMDD_COT全履歴.csv」（UTF-8 BOM付き）"],
        [""],
        ["データ種類"],
        ["先物のみ = Futures Only / 先物+オプション = Futures and Options Combined（オプションはデルタ換算）"],
        ["Supplemental CIT は CFTC が先物+オプション版のみ公表"],
    ]
    return pd.DataFrame(lines)


def style_sheet(worksheet, frame, index_column_name=None):
    """見出し色・列幅・数値書式・フィルタ・固定表示を設定する。"""
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    integer_columns = ["ロング", "ショート", "ネット", "前週比", "総建玉"]
    for position, name in enumerate(frame.columns, start=1):
        letter = get_column_letter(position)
        width = max(10, min(28, int(frame[name].astype(str).str.len().max() * 1.6) if len(frame) else 10))
        worksheet.column_dimensions[letter].width = max(width, len(name) * 2 + 2)
        if name in integer_columns:
            for cell in worksheet[letter][1:]:
                cell.number_format = "#,##0;[Red]-#,##0"
        if name in ["建玉比(%)", "COT指数"]:
            for cell in worksheet[letter][1:]:
                cell.number_format = "0.0"

    if index_column_name and len(frame):
        letter = get_column_letter(list(frame.columns).index(index_column_name) + 1)
        cell_range = f"{letter}2:{letter}{len(frame) + 1}"
        rule = ColorScaleRule(
            start_type="num", start_value=0, start_color="F8696B",
            mid_type="num", mid_value=50, mid_color="FFFFFF",
            end_type="num", end_value=100, end_color="63BE7B",
        )
        worksheet.conditional_formatting.add(cell_range, rule)


def write_excel(metrics, markets, path, generated_at):
    """Excelファイルを書き出す。"""
    labeled = attach_labels(metrics, markets)
    latest_mask = labeled["date"] == labeled.groupby(["report", "variant"])["date"].transform("max")
    latest = labeled[latest_mask]
    latest_date = latest["date"].max()

    readme = build_readme(latest_date, generated_at)
    summary_fo = build_summary(latest, VARIANT_FUTURES_ONLY)
    summary_combined = build_summary(latest, VARIANT_COMBINED)
    history_tff = build_history(labeled, "TFF", SPECULATOR_CATEGORIES["TFF"])
    history_disagg = build_history(labeled, "Disagg", SPECULATOR_CATEGORIES["Disagg"])
    history_legacy = build_history(labeled, "Legacy", SPECULATOR_CATEGORIES["Legacy"])
    history_cit = build_history(labeled, "CIT", SPECULATOR_CATEGORIES["CIT"])

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="説明", index=False, header=False)
        summary_fo.to_excel(writer, sheet_name="サマリー_先物のみ", index=False)
        summary_combined.to_excel(writer, sheet_name="サマリー_先物OP込み", index=False)
        history_tff.to_excel(writer, sheet_name="履歴_TFF", index=False)
        history_disagg.to_excel(writer, sheet_name="履歴_Disagg", index=False)
        history_legacy.to_excel(writer, sheet_name="履歴_Legacy", index=False)
        history_cit.to_excel(writer, sheet_name="履歴_CIT", index=False)

        writer.sheets["説明"].column_dimensions["A"].width = 110
        writer.sheets["説明"]["A1"].font = Font(bold=True, size=14)
        style_sheet(writer.sheets["サマリー_先物のみ"], summary_fo, "COT指数")
        style_sheet(writer.sheets["サマリー_先物OP込み"], summary_combined, "COT指数")
        style_sheet(writer.sheets["履歴_TFF"], history_tff)
        style_sheet(writer.sheets["履歴_Disagg"], history_disagg)
        style_sheet(writer.sheets["履歴_Legacy"], history_legacy)
        style_sheet(writer.sheets["履歴_CIT"], history_cit)
    return latest_date


def write_history_csv(metrics, markets, path):
    """全レポート・全区分の履歴をCSVで書き出す。

    Excelで直接開いても文字化けしないよう、UTF-8（BOM付き）で保存する。
    （Shift_JISは銘柄名の記号で変換エラーが起きうるため不採用）
    """
    labeled = attach_labels(metrics, markets)
    table = build_history(labeled)
    table.to_csv(path, index=False, encoding="utf-8-sig")
